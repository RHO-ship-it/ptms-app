"""Tests de bout en bout : carte Grubbs détaillée, choix des estimateurs du rapport, rapport Word, kappa de Fleiss.
Lancer :  python tests/e2e_features.py   (nécessite les dépendances de requirements.txt)"""
import os, sys, io, contextlib, tempfile
DB=os.path.join(tempfile.mkdtemp(), "ptms_test.db")
os.environ.update(DATABASE_URL=f"sqlite:///{DB}", SECRET_KEY="k")
if os.path.exists(DB): os.remove(DB)
ROOT=os.path.abspath(os.path.join(os.path.dirname(__file__),".."))
sys.path.insert(0,ROOT); os.chdir(ROOT)
from flask import template_rendered
from app import create_app
from extensions import db
from models import *
app=create_app(); R=[]
def chk(n,c,d=""): R.append((bool(c),n,d)); 
VALS=[290,310,350,380,395,420,445]
with contextlib.redirect_stdout(io.StringIO()):
  with app.app_context():
    a=User(username="admin",email="a@x",full_name="A",role="admin",active=True); a.set_password("pw"); db.session.add(a)
    c=Campaign(reference="C1",name="C1",status="active"); db.session.add(c); db.session.flush()
    t=Trial(campaign_id=c.id,name="T1",trial_type="quantitative",unit="N",sigma_pt_method="robuste")
    q1=Trial(campaign_id=c.id,name="Q1",trial_type="qualitative"); q2=Trial(campaign_id=c.id,name="Q2",trial_type="qualitative")
    db.session.add_all([t,q1,q2]); db.session.flush()
    for i,v in enumerate(VALS):
        lab=Laboratory(code=f"L{i+1}",name=f"Labo {i+1}",active=True); db.session.add(lab); db.session.flush()
        db.session.add(CampaignLaboratory(campaign_id=c.id,laboratory_id=lab.id))
        for tr,kw in ((t,{"value":v}),(q1,{"qualitative_result":"Conforme" if i<5 else "Non conforme"}),(q2,{"qualitative_result":"Conforme" if i%2==0 else "Non conforme"})):
            f=ParticipantForm(trial_id=tr.id,laboratory_id=lab.id,status="validated"); db.session.add(f); db.session.flush()
            db.session.add(Submission(form_id=f.id,validated=True,**kw))
    db.session.commit(); TID,Q1,CID=t.id,q1.id,c.id
cl=app.test_client(); cl.post("/auth/login",data={"username":"admin","password":"pw"})
with contextlib.redirect_stdout(io.StringIO()):
    r=cl.get(f"/traitement/atelier?trial_id={TID}")
html=r.get_data(as_text=True)
chk("Atelier quantitatif 200",r.status_code==200)
chk("Carte Grubbs détaillé présente","Test de Grubbs détaillé" in html and "Itération n° 1" in html)
chk("Valeurs Excel affichées (1.324 / 1.412 / 2.139 / 2.020 / 19250.000 / 8000.000 / 5330.000)",all(s in html for s in ("1.324","1.412","2.139","2.020","19250.000","8000.000","5330.000","0.416","0.277","0.031","0.071")))
chk("Liste déroulante des estimateurs présente","Estimateurs affichés dans le rapport" in html and 'name="report_estimators"' in html)
with contextlib.redirect_stdout(io.StringIO()):
    r=cl.post(f"/traitement/report-estimators/{TID}",data={"report_estimators":["algorithm_a","classical","bidon"]})
with app.app_context():
    tr=db.session.get(Trial,TID); chk("Sélection enregistrée (clés invalides ignorées)",tr.selected_report_estimators()==["algorithm_a","classical"],str(tr.report_estimators))
    chk("Journal d'audit",AuditLog.query.filter_by(action="set_report_estimators").count()==1)
with contextlib.redirect_stdout(io.StringIO()):
    r=cl.get(f"/reports/campaign/{CID}/export-report-word")
chk("Rapport Word complet 200",r.status_code==200)
import docx
d=docx.Document(io.BytesIO(r.data)); txt="\n".join(p.text for p in d.paragraphs)
cells=" ".join(c.text for tb in d.tables for row in tb.rows for c in row.cells)
chk("Rapport : tableau des estimateurs choisis","Estimateurs retenus pour ce rapport" in txt and "Algorithme A (ISO 13528)" in cells and "Moyenne / écart-type classiques" in cells and "Médiane / MADe" not in cells)
chk("Rapport : texte d'introduction mis à jour","x_pt" in txt and "choisis par l'administrateur" in txt)
with contextlib.redirect_stdout(io.StringIO()):
    r=cl.get(f"/traitement/atelier?trial_id={Q1}")
h2=r.get_data(as_text=True)
chk("Atelier qualitatif (kappa de Fleiss sur 2 essais) 200",r.status_code==200 and "κ de Fleiss" in h2)
# sans sélection -> pas de tableau
with contextlib.redirect_stdout(io.StringIO()):
    cl.post(f"/traitement/report-estimators/{TID}",data={})
    r=cl.get(f"/reports/campaign/{CID}/export-report-word")
d=docx.Document(io.BytesIO(r.data)); txt="\n".join(p.text for p in d.paragraphs)
chk("Sans sélection : pas de tableau supplémentaire","Estimateurs retenus pour ce rapport" not in txt)
# rapport publié -> choix bloqué
with app.app_context():
    cc=db.session.get(Campaign,CID); cc.status="closed"; cc.report_published=True; db.session.commit()
with contextlib.redirect_stdout(io.StringIO()):
    cl.post(f"/traitement/report-estimators/{TID}",data={"report_estimators":["median_made"]})
with app.app_context(): chk("Rapport publié : choix bloqué",db.session.get(Trial,TID).selected_report_estimators()==[])
for ok,n,dd in R: print("OK  " if ok else "FAIL",n,dd)
bad=[n for o,n,_ in R if not o]
print(f"=== {len(R)-len(bad)}/{len(R)} OK ===")
sys.exit(1 if bad else 0)
